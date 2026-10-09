"""管理扩展：永久删除、批量操作和实际文档预览。"""
import io
import os
import tempfile
import unittest
from unittest.mock import patch
from fastapi.testclient import TestClient
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from tests.test_service import FixtureEmbedding
from app.main import create_app
from app.storage import Storage


class ManagementTests(unittest.TestCase):
    def setUp(self):
        environment = patch.dict(os.environ, {'RAG_ADMIN_TOKEN': 'admin', 'RAG_MCP_TOKEN': ''})
        environment.start()
        self.addCleanup(environment.stop)
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.application = create_app(self.directory.name, FixtureEmbedding())
        self.client = TestClient(self.application, base_url='http://127.0.0.1', headers={'Authorization': 'Bearer admin'})
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)
        self.client.put('/api/config', json={'base_url': 'https://example.org/v1', 'model': 'fixture'})

    def upload(self, name='说明.txt', content=b'original', library='default'):
        response = self.client.post('/api/documents', params={'library_id': library}, files={'file': (name, content)})
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()

    def test_delete_library_cascades_and_restart_does_not_recreate(self):
        folder = self.client.post('/api/folders', json={'library_id': 'default', 'name': '一级'}).json()['id']
        self.client.post('/api/folders', json={'library_id': 'default', 'parent_id': folder, 'name': '二级'})
        doc = self.upload()
        self.client.patch('/api/documents/' + doc['id'] + '/location', json={'library_id': 'default', 'folder_id': folder})
        response = self.client.patch('/api/libraries/default', json={'name': '我的资料'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['name'], '我的资料')
        self.assertEqual(self.client.delete('/api/libraries/default').status_code, 200)
        store = Storage(self.directory.name)
        with store.connect() as db:
            self.assertFalse(db.execute("SELECT 1 FROM libraries WHERE id='default'").fetchone())
            self.assertFalse(db.execute('SELECT 1 FROM documents').fetchone())
            self.assertFalse(db.execute('SELECT 1 FROM chunks').fetchone())
            self.assertFalse(db.execute('SELECT 1 FROM folders').fetchone())
            self.assertFalse(db.execute('PRAGMA foreign_key_check').fetchall())
        self.assertFalse(list(store.uploads.iterdir()))
        for lib in self.client.get('/api/libraries').json()['libraries']:
            self.assertEqual(self.client.delete('/api/libraries/' + lib['id']).status_code, 200)
        Storage(self.directory.name)
        self.assertEqual(self.client.get('/api/libraries').json()['libraries'], [])
        self.assertEqual(self.client.post('/api/documents', files={'file': ('a.txt', b'a')}).status_code, 400)

    def test_batch_move_validates_all_before_changes_and_preserves_vectors(self):
        a, b = self.upload('a.txt'), self.upload('b.txt')
        folder = self.client.post('/api/folders', json={'library_id': 'faq', 'name': '培训'}).json()['id']
        with self.application.state.service.storage.connect() as db:
            before = [tuple(r) for r in db.execute('SELECT * FROM chunks ORDER BY document_id,chunk_index')]
        body = {'document_ids': [a['id'], 'missing'], 'library_id': 'faq', 'folder_id': folder}
        self.assertEqual(self.client.post('/api/documents/batch/move', json=body).status_code, 404)
        self.assertTrue(all(d['library_id'] == 'default' for d in self.client.get('/api/documents').json()['documents']))
        body['document_ids'] = [a['id'], b['id']]
        self.assertEqual(self.client.post('/api/documents/batch/move', json=body).status_code, 200)
        self.assertTrue(all(d['folder_id'] == folder for d in self.client.get('/api/documents').json()['documents']))
        with self.application.state.service.storage.connect() as db:
            self.assertEqual([tuple(r) for r in db.execute('SELECT * FROM chunks ORDER BY document_id,chunk_index')], before)
        self.assertEqual(self.client.post('/api/documents/batch/delete', json={'document_ids': [a['id'], 'missing']}).status_code, 404)
        self.assertEqual(len(self.client.get('/api/documents').json()['documents']), 2)
        self.assertEqual(self.client.post('/api/documents/batch/delete', json={'document_ids': [a['id'], b['id']]}).status_code, 200)
        self.assertEqual(self.client.get('/api/documents').json()['documents'], [])

    def test_deleted_last_key_still_protects_mcp_and_revoked_key_can_be_removed(self):
        key = self.client.post('/api/keys', json={'name': '测试', 'all_libraries': True}).json()
        self.client.delete('/api/keys/' + key['id'])
        self.assertEqual(self.client.delete('/api/keys/' + key['id'] + '/permanent').status_code, 200)
        self.assertEqual(self.client.get('/api/keys').json()['keys'], [])
        self.assertEqual(self.client.get('/mcp/', headers={'Authorization': ''}).status_code, 401)
        self.assertEqual(self.client.get('/api/documents', headers={'Authorization': 'Bearer ' + key['key']}).status_code, 401)

    def test_readonly_key_cannot_batch_mutate_or_delete_library(self):
        item = self.upload()
        key = self.client.post('/api/keys', json={'name': '只读', 'all_libraries': True}).json()['key']
        headers = {'Authorization': 'Bearer ' + key}
        for endpoint, payload in [('/api/documents/batch/delete', {'document_ids': [item['id']]}), ('/api/documents/batch/move', {'document_ids': [item['id']], 'library_id': 'faq'})]:
            self.assertEqual(self.client.post(endpoint, headers=headers, json=payload).status_code, 401)
        self.assertEqual(self.client.delete('/api/libraries/default', headers=headers).status_code, 401)
        self.assertEqual(len(self.client.get('/api/documents').json()['documents']), 1)

    def test_permanent_key_delete_invalidates_previously_issued_download(self):
        from urllib.parse import urlsplit
        self.upload()
        key = self.client.post('/api/keys', json={'name': '待删除', 'all_libraries': True}).json()
        url = urlsplit(self.client.get('/api/documents', headers={'Authorization': 'Bearer ' + key['key']}).json()['documents'][0]['download_url'])
        endpoint = url.path + '?' + url.query
        self.assertEqual(self.client.get(endpoint, headers={'Authorization': ''}).status_code, 200)
        self.assertEqual(self.client.delete('/api/keys/' + key['id'] + '/permanent').status_code, 200)
        self.assertEqual(self.client.get(endpoint, headers={'Authorization': ''}).status_code, 401)

    def test_preview_formats_are_readable_and_cannot_execute_document_markup(self):
        fixtures = [('a.txt', b'<script>alert(1)</script>hello', '&lt;script&gt;'), ('a.md', b'# Heading\n\n<script>alert(1)</script>', '<h1>Heading</h1>'), ('a.csv', b'name,value\nradar,42', '<td>42</td>')]
        doc = Document(); doc.add_paragraph('Word body'); doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Word cell'
        stream = io.BytesIO(); doc.save(stream); fixtures.append(('a.docx', stream.getvalue(), 'Word cell'))
        book = Workbook(); book.active['A1'] = 'Excel cell'; stream = io.BytesIO(); book.save(stream); fixtures.append(('a.xlsx', stream.getvalue(), 'Excel cell'))
        slides = Presentation(); slide = slides.slides.add_slide(slides.slide_layouts[1]); slide.shapes.title.text = 'PPT title'; stream = io.BytesIO(); slides.save(stream); fixtures.append(('a.pptx', stream.getvalue(), 'PPT title'))
        for name, content, expected in fixtures:
            with self.subTest(name=name):
                item = self.upload(name, content)
                self.client.patch('/api/documents/' + item['id'], json={'enabled': False})
                response = self.client.get('/api/documents/' + item['id'] + '/preview')
                self.assertEqual(response.status_code, 200, response.text)
                self.assertIn(expected, response.text)
                self.assertNotIn('<script>', response.text)
        self.client.put('/api/service', json={'enabled': False})
        self.assertEqual(self.client.get('/api/documents/' + item['id'] + '/preview').status_code, 200)
        key = self.client.post('/api/keys', json={'name': '只读', 'all_libraries': True}).json()['key']
        self.assertEqual(self.client.get('/api/documents/' + item['id'] + '/preview', headers={'Authorization': 'Bearer ' + key}).status_code, 401)

    def test_pdf_preview_returns_original_bytes_and_admin_text_works_when_paused(self):
        from tests.document_fixtures import text_pdf
        raw = text_pdf()
        item = self.upload('manual.pdf', raw)
        self.client.put('/api/service', json={'enabled': False})
        response = self.client.get('/api/documents/' + item['id'] + '/preview')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['content-type'], 'application/pdf')
        self.assertEqual(response.content, raw)
        self.assertEqual(self.client.get('/api/documents/' + item['id'] + '/admin-content').json()['text'], 'PDF preview test')
