import io
import unittest

import openpyxl
import pptx
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.documents import extract_text, split_text


class DocumentTests(unittest.TestCase):
    def test_text_pdf_extracts_visible_text(self):
        pdf = PdfWriter()
        page = pdf.add_blank_page(width=200, height=200)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 150 Td (Check network before installation.) Tj ET")
        page[NameObject("/Contents")] = stream
        output = io.BytesIO()
        pdf.write(output)
        _, text = extract_text("manual.pdf", output.getvalue())
        self.assertEqual(text, "Check network before installation.")

    def test_docx_extracts_paragraphs_and_table_cells(self):
        document = Document()
        document.add_paragraph("设备安装说明")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "接口"
        table.cell(0, 1).text = "以太网"
        output = io.BytesIO()
        document.save(output)
        suffix, text = extract_text("设备.docx", output.getvalue())
        self.assertEqual(suffix, ".docx")
        self.assertEqual(text, "设备安装说明\n接口 | 以太网")

    def test_xlsx_extracts_sheets_and_table(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "调试记录"
        sheet.append(["模块", "状态", "备注"])
        sheet.append(["电源", "正常", "电压稳定"])
        output = io.BytesIO()
        workbook.save(output)
        suffix, text = extract_text("工程记录.xlsx", output.getvalue())
        self.assertEqual(suffix, ".xlsx")
        self.assertIn("调试记录", text)
        self.assertIn("电源", text)
        self.assertIn("电压稳定", text)

    def test_pptx_extracts_slides_and_text(self):
        presentation = pptx.Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = "产品架构培训"
        output = io.BytesIO()
        presentation.save(output)
        suffix, text = extract_text("培训.pptx", output.getvalue())
        self.assertEqual(suffix, ".pptx")
        self.assertIn("产品架构培训", text)

    def test_csv_extracts_utf8_and_gbk(self):
        utf8_content = "项目,数值\n传感器A,100".encode("utf-8")
        suffix, text = extract_text("data_utf8.csv", utf8_content)
        self.assertEqual(suffix, ".csv")
        self.assertIn("传感器A", text)

        gbk_content = "项目,数值\n传感器B,200".encode("gbk")
        suffix, text = extract_text("data_gbk.csv", gbk_content)
        self.assertEqual(suffix, ".csv")
        self.assertIn("传感器B", text)

    def test_blank_pdf_rejected_with_ocr_direction(self):
        pdf = PdfWriter()
        pdf.add_blank_page(width=200, height=200)
        output = io.BytesIO()
        pdf.write(output)
        with self.assertRaisesRegex(ValueError, "扫描版 PDF"):
            extract_text("scan.pdf", output.getvalue())

    def test_utf8_bom_and_long_text_chunk_overlap(self):
        _, text = extract_text("note.md", b"\xef\xbb\xbf" + "说明文字".encode())
        self.assertEqual(text, "说明文字")
        chunks = split_text("甲" * 680 + "乙" * 120 + "丙" * 200)
        self.assertEqual(chunks, ["甲" * 680 + "乙" * 120, "乙" * 120 + "丙" * 200])

    def test_invalid_encoding_reports_utf8(self):
        with self.assertRaisesRegex(ValueError, "UTF-8"):
            extract_text("note.txt", b"\xff\xfe\x00")

    def test_unsupported_extension_rejected(self):
        with self.assertRaisesRegex(ValueError, "仅支持"):
            extract_text("app.exe", b"binary content")
