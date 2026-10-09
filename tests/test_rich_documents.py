"""图片参与解析、保留来源，以及工作表和幻灯片的可视结构。"""
import io
import tempfile
import unittest
from pathlib import Path
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.drawing.image import Image as SheetImage
from pptx import Presentation
from pptx.util import Inches as SlideInches
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, NumberObject, DecodedStreamObject
from app import documents
from app.preview import render_preview


def image_bytes():
    image = Image.new('RGB', (400, 100), 'white')
    ImageDraw.Draw(image).text((20, 30), 'OCR image fixture', fill='black')
    stream = io.BytesIO(); image.save(stream, 'PNG'); return stream.getvalue()


def scanned_pdf():
    pdf = PdfWriter(); page = pdf.add_blank_page(width=400, height=100)
    image = Image.open(io.BytesIO(image_bytes())).convert('RGB')
    obj = DecodedStreamObject(); obj.set_data(image.tobytes())
    obj.update({NameObject('/Type'):NameObject('/XObject'), NameObject('/Subtype'):NameObject('/Image'), NameObject('/Width'):NumberObject(400), NameObject('/Height'):NumberObject(100), NameObject('/ColorSpace'):NameObject('/DeviceRGB'), NameObject('/BitsPerComponent'):NumberObject(8)})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/XObject'):DictionaryObject({NameObject('/Im1'):pdf._add_object(obj)})})
    stream = DecodedStreamObject(); stream.set_data(b'q 400 0 0 100 0 0 cm /Im1 Do Q'); page[NameObject('/Contents')] = stream
    output = io.BytesIO(); pdf.write(output); return output.getvalue()


def office_fixture(suffix):
    output = io.BytesIO()
    if suffix == '.docx':
        doc = Document(); doc.add_paragraph('Native Word text'); doc.add_picture(io.BytesIO(image_bytes()), width=Inches(3)); doc.save(output)
    elif suffix == '.xlsx':
        book = Workbook(); sheet = book.active; sheet.title = '样式工作表'; sheet['A1']='Styled cell'; sheet.merge_cells('A1:C1'); sheet['A1'].fill=PatternFill('solid', fgColor='FFCC00'); sheet['A1'].font=Font(bold=True)
        sheet.add_image(SheetImage(io.BytesIO(image_bytes())), 'B3'); book.create_sheet('第二表')['A1']='Second worksheet'; book.save(output)
    else:
        slides = Presentation(); slide = slides.slides.add_slide(slides.slide_layouts[6]); box = slide.shapes.add_textbox(SlideInches(1),SlideInches(1),SlideInches(4),SlideInches(1)); box.text='Native slide text'
        slide.shapes.add_picture(io.BytesIO(image_bytes()), SlideInches(1), SlideInches(2), width=SlideInches(4)); slides.save(output)
    return output.getvalue()


class RichDocumentTests(unittest.TestCase):
    def extract(self, name, content, recognizer):
        method = getattr(documents, 'extract_document', None)
        self.assertTrue(callable(method), '需要能提取图片 OCR 和来源的解析入口')
        return method(name, content, recognizer=recognizer)

    def test_scanned_pdf_ocr_has_page_provenance(self):
        suffix, text, info = self.extract('scan.pdf', scanned_pdf(), lambda raw: 'Image radar 123')
        self.assertEqual(suffix, '.pdf'); self.assertIn('Image radar 123', text); self.assertIn('OCR · PDF 第 1 页', text)
        self.assertEqual(info['recognized'], 1); self.assertEqual(info['state'], 'complete')

    def test_office_embedded_images_are_indexable_and_marked(self):
        for suffix in ['.docx', '.xlsx', '.pptx']:
            with self.subTest(suffix=suffix):
                _, text, info = self.extract('fixture'+suffix, office_fixture(suffix), lambda raw: 'Image radar 123')
                self.assertIn('Image radar 123', text); self.assertIn('OCR · ', text); self.assertEqual(info['recognized'], 1)
                self.assertIn('Native' if suffix != '.xlsx' else 'Styled cell', text)

    def test_failed_ocr_is_visible_and_keeps_native_text(self):
        def failed(raw): raise ValueError('识别超时')
        _, text, info = self.extract('fixture.docx', office_fixture('.docx'), failed)
        self.assertIn('Native Word text', text); self.assertEqual(info['state'], 'partial'); self.assertTrue(info['warnings'])

    def test_spreadsheet_preserves_tabs_merge_fill_and_picture(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'fixture.xlsx'; path.write_bytes(office_fixture('.xlsx'))
            html = render_preview({'suffix':'.xlsx','text':''}, path)
            self.assertIn('data-sheet-tab=', html); self.assertIn('colspan="3"', html); self.assertIn('#FFCC00', html); self.assertIn('data:image/png;base64,', html)

    def test_presentation_preserves_slide_coordinates_and_picture(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'fixture.pptx'; path.write_bytes(office_fixture('.pptx'))
            html=render_preview({'suffix':'.pptx','text':''},path)
            self.assertIn('<svg', html); self.assertIn('viewBox=', html); self.assertIn('Native slide text',html); self.assertIn('data:image/png;base64,',html)
