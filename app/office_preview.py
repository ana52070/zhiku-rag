"""轻量只读 Office 视图：保留表格样式、图片与幻灯片坐标。"""
import base64
import io
from html import escape
from datetime import date, datetime
from PIL import Image
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN


def image_url(raw):
    try:
        with Image.open(io.BytesIO(raw)) as image:
            if image.width * image.height > 40_000_000:
                return None
            image.thumbnail((1800, 1800))
            output=io.BytesIO(); image.save(output, 'PNG')
            return 'data:image/png;base64,' + base64.b64encode(output.getvalue()).decode()
    except (OSError, ValueError, Image.DecompressionBombError):
        return None


def color(value, fallback='#ffffff'):
    if value is not None and getattr(value, 'type', None) == 'rgb':
        raw = str(value.rgb)[-6:]
        if len(raw) == 6 and all(c in '0123456789abcdefABCDEF' for c in raw): return '#' + raw
    return fallback


def cell_style(cell):
    parts = []
    if cell.fill.patternType == 'solid': parts.append('background:'+color(cell.fill.fgColor))
    font=cell.font
    if font.bold: parts.append('font-weight:700')
    if font.italic: parts.append('font-style:italic')
    parts.append('color:'+color(font.color,'#182b41'))
    if font.sz: parts.append('font-size:'+str(min(float(font.sz),60))+'pt')
    alignment=cell.alignment
    if alignment.horizontal in {'left','right','center','justify'}: parts.append('text-align:'+alignment.horizontal)
    if alignment.vertical in {'top','center','bottom'}: parts.append('vertical-align:'+('middle' if alignment.vertical=='center' else alignment.vertical))
    if alignment.wrap_text: parts.append('white-space:pre-wrap')
    for side in ['top','right','bottom','left']:
        border=getattr(cell.border,side)
        if border and border.style: parts.append('border-'+side+':1px solid '+color(border.color,'#4f6175'))
    return ';'.join(parts)


def cell_value(cell):
    value=cell.value
    if value is None: return ''
    if isinstance(value,(datetime,date)): return value.isoformat(sep=' ') if isinstance(value,datetime) else value.isoformat()
    if isinstance(value,(int,float)) and not isinstance(value,bool):
        fmt=cell.number_format
        if '%' in fmt:
            decimals=len(fmt.split('.')[1].split('%')[0]) if '.' in fmt else 0
            return f'{value * 100:.{min(decimals,8)}f}%'
        if '.' in fmt and any(char in fmt for char in '0#'):
            import re
            match=re.search(r'\.([0#]+)',fmt)
            if match: return f'{value:,.{min(len(match[1]),8)}f}' if ',' in fmt else f'{value:.{min(len(match[1]),8)}f}'
    return str(value)


def spreadsheet_html(path):
    book=load_workbook(path,data_only=True)
    try:
        tabs=[]; panels=[]
        for index,sheet in enumerate(book):
            if sheet.sheet_state != 'visible': continue
            selected=not tabs
            tabs.append(f'<button class="sheet-tab" data-sheet-tab="{index}" aria-selected="{str(selected).lower()}">{escape(sheet.title)}</button>')
            rows=min(max(sheet.max_row,20),2000); columns=min(max(sheet.max_column,10),150)
            widths=[min(max(float(sheet.column_dimensions[get_column_letter(column)].width or 13)*7+5,32),600) for column in range(1,columns+1)]
            heights=[min(max(float(sheet.row_dimensions[row].height or 15)*96/72,20),600) for row in range(1,rows+1)]
            merged={}; skipped=set()
            for area in sheet.merged_cells.ranges:
                if area.min_row > rows or area.min_col > columns: continue
                merged[(area.min_row,area.min_col)]=(min(area.max_row,rows)-area.min_row+1,min(area.max_col,columns)-area.min_col+1)
                for row in range(area.min_row,min(area.max_row,rows)+1):
                    for column in range(area.min_col,min(area.max_col,columns)+1):
                        if (row,column)!=(area.min_row,area.min_col): skipped.add((row,column))
            header='<tr><th class="sheet-corner"></th>'+''.join('<th>'+get_column_letter(column)+'</th>' for column in range(1,columns+1))+'</tr>'
            body=[]
            for row in range(1,rows+1):
                cells=[f'<th class="sheet-row-number">{row}</th>']
                for column in range(1,columns+1):
                    if (row,column) in skipped: continue
                    cell=sheet.cell(row,column); span=merged.get((row,column),(1,1))
                    cells.append(f'<td rowspan="{span[0]}" colspan="{span[1]}" style="{cell_style(cell)}">'+escape(cell_value(cell))+'</td>')
                body.append(f'<tr style="height:{heights[row-1]:.1f}px">'+''.join(cells)+'</tr>')
            images=[]
            for image in sheet._images:
                uri=image_url(image._data())
                if not uri: images.append('<p class="preview-warning">一张图片格式无法显示，请下载源文件查看。</p>'); continue
                anchor=image.anchor
                start=getattr(anchor,'_from',None)
                if not start: continue
                x=42+sum(widths[:start.col])+start.colOff/9525
                y=28+sum(heights[:start.row])+start.rowOff/9525
                ext=getattr(anchor,'ext',None)
                width=ext.cx/9525 if ext else image.width; height=ext.cy/9525 if ext else image.height
                end=getattr(anchor,'to',None)
                if end:
                    width=max(1,42+sum(widths[:end.col])+end.colOff/9525-x)
                    height=max(1,28+sum(heights[:end.row])+end.rowOff/9525-y)
                images.append(f'<img class="sheet-image" alt="工作表内嵌图片" src="{uri}" style="left:{x:.1f}px;top:{y:.1f}px;width:{width:.1f}px;height:{height:.1f}px">')
            cols='<col style="width:42px">'+''.join(f'<col style="width:{width:.1f}px">' for width in widths)
            note='<p class="preview-warning">工作表较大，预览显示前 2000 行、150 列；完整内容请下载源文件。</p>' if sheet.max_row>2000 or sheet.max_column>150 else ''
            panels.append(f'<section data-sheet-panel="{index}"'+('' if selected else ' hidden')+'>'+note+'<div class="sheet-scroll"><div class="sheet-surface"><table class="spreadsheet-grid"><colgroup>'+cols+'</colgroup><thead>'+header+'</thead><tbody>'+''.join(body)+'</tbody></table>'+''.join(images)+'</div></div></section>')
        return '<div class="spreadsheet-view"><div class="sheet-tabs" role="tablist">'+''.join(tabs)+'</div>'+''.join(panels)+'</div>'
    finally:
        book.close()


def _ppt_color(color_format, fallback):
    try:
        rgb=color_format.rgb
        if rgb is not None: return '#'+str(rgb)
    except (AttributeError,ValueError): pass
    return fallback


def _shape_svg(shape):
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        transform=shape._element.grpSpPr.xfrm
        scale_x=transform.ext.cx / max(transform.chExt.cx,1); scale_y=transform.ext.cy / max(transform.chExt.cy,1)
        return f'<g transform="translate({transform.off.x} {transform.off.y}) scale({scale_x} {scale_y}) translate({-transform.chOff.x} {-transform.chOff.y})">'+''.join(_shape_svg(item) for item in shape.shapes)+'</g>'
    x,y,w,h=shape.left,shape.top,shape.width,shape.height
    parts=[]
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        uri=image_url(shape.image.blob)
        if uri: return f'<image x="{x}" y="{y}" width="{w}" height="{h}" href="{uri}" preserveAspectRatio="none"/>'
        return f'<text x="{x}" y="{y+180000}" font-size="160000">图片格式暂不支持</text>'
    try:
        if shape.fill.type is not None:
            fill=_ppt_color(shape.fill.fore_color,'#ffffff')
            parts.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="{fill}"/>')
    except (AttributeError,TypeError,ValueError): pass
    if shape.has_table:
        table=shape.table; tx=x; ty=y
        for row in table.rows:
            tx=x
            for index,cell in enumerate(row.cells):
                width=table.columns[index].width
                parts.append(f'<rect x="{tx}" y="{ty}" width="{width}" height="{row.height}" fill="#ffffff" stroke="#b5c4d5" stroke-width="9000"/>')
                parts.append(f'<text x="{tx+70000}" y="{ty+min(row.height,240000)}" font-size="160000">'+escape(cell.text)+'</text>')
                tx+=width
            ty+=row.height
    if shape.has_text_frame:
        import textwrap
        frame=shape.text_frame; cursor=y+frame.margin_top
        for paragraph in frame.paragraphs:
            runs=paragraph.runs
            font=runs[0].font if runs else paragraph.font
            size=font.size or paragraph.font.size or 228600
            fg=_ppt_color(font.color,'#172d46')
            text=paragraph.text
            lines=textwrap.wrap(text,width=max(int((w-frame.margin_left-frame.margin_right)/max(size*.55,1)),1),replace_whitespace=False) or ['']
            for line in lines:
                cursor+=size*1.25
                alignment=paragraph.alignment
                position=x+w/2 if alignment==PP_ALIGN.CENTER else x+w-frame.margin_right if alignment==PP_ALIGN.RIGHT else x+frame.margin_left
                anchor='middle' if alignment==PP_ALIGN.CENTER else 'end' if alignment==PP_ALIGN.RIGHT else 'start'
                parts.append(f'<text x="{position}" y="{cursor}" font-size="{size}" fill="{fg}" text-anchor="{anchor}" font-weight="{700 if font.bold else 400}">'+escape(line)+'</text>')
    return ''.join(parts)


def presentation_html(path):
    presentation=Presentation(str(path))
    parts=['<div class="presentation-view">']
    for number,slide in enumerate(presentation.slides,1):
        background='#ffffff'
        try:
            if slide.background.fill.type is not None: background=_ppt_color(slide.background.fill.fore_color,background)
        except (AttributeError,TypeError,ValueError): pass
        parts.append(f'<section class="slide-page"><h3>第 {number} 页</h3><svg role="img" aria-label="第 {number} 页幻灯片" viewBox="0 0 {presentation.slide_width} {presentation.slide_height}"><rect width="100%" height="100%" fill="{background}"/>'+''.join(_shape_svg(shape) for shape in slide.shapes)+'</svg></section>')
    return ''.join(parts)+'</div>'
