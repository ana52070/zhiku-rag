import io
import zipfile
from pathlib import Path

import openpyxl
from docx import Document
from markitdown import MarkItDown
from pypdf import PdfReader

MAX_UPLOAD = 50 * 1024 * 1024
CHUNK_SIZE = 800
CHUNK_OVERLAP = 120
MAX_CHUNKS = 10000

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".xlsx", ".pptx", ".csv"}


def split_text(text):
    return [text[start:start + CHUNK_SIZE] for start in range(0, len(text), CHUNK_SIZE - CHUNK_OVERLAP) if text[start:start + CHUNK_SIZE].strip()]


def _extract_xlsx(content):
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    blocks = []
    for name in workbook.sheetnames:
        sheet = workbook[name]
        raw_rows = []
        max_cols = 0
        for row in sheet.iter_rows(values_only=True):
            if not any(cell is not None and str(cell).strip() for cell in row):
                continue
            cells = [str(cell).replace("\n", " ").strip() if cell is not None else "" for cell in row]
            while cells and not cells[-1]:
                cells.pop()
            if cells:
                max_cols = max(max_cols, len(cells))
                raw_rows.append(cells)
        if raw_rows:
            formatted_rows = []
            for r in raw_rows:
                padded = r + [""] * (max_cols - len(r))
                formatted_rows.append("| " + " | ".join(padded) + " |")
            sep = "| " + " | ".join(["---"] * max_cols) + " |"
            table = formatted_rows[0] + "\n" + sep + "\n" + "\n".join(formatted_rows[1:]) if len(formatted_rows) > 1 else formatted_rows[0]
            blocks.append(f"## {name}\n\n{table}")
    workbook.close()
    return "\n\n".join(blocks)


def extract_text(filename, content, allow_empty=False):
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError("仅支持 TXT、Markdown、文字 PDF、Word (DOCX)、Excel (XLSX)、PPT (PPTX) 和 CSV 文件。")
    if len(content) > MAX_UPLOAD:
        raise ValueError("文件超过 50 MiB，请拆分后上传。")
    try:
        if suffix in {".txt", ".md"}:
            text = content.decode("utf-8-sig")
        elif suffix == ".csv":
            decoded = None
            for enc in ("utf-8-sig", "gb18030", "gbk"):
                try:
                    decoded = content.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            if decoded is None:
                raise ValueError("CSV 编码无法识别，请另存为 UTF-8 或 GBK 后上传。")
            md = MarkItDown()
            result = md.convert_stream(io.BytesIO(decoded.encode("utf-8")), file_extension=".csv")
            text = result.text_content
        elif suffix == ".pdf":
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise ValueError("请先解除 PDF 密码保护。")
            if len(reader.pages) > 1000:
                raise ValueError("PDF 超过 1000 页，请拆分后上传。")
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 200 * 1024 * 1024:
                    raise ValueError("DOCX 解压后超过 200 MiB，请拆分后上传。")
            document = Document(io.BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                paragraphs.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            text = "\n".join(paragraphs)
            from xml.etree import ElementTree as ET
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                root = ET.fromstring(archive.read('word/document.xml'))
                for textbox in root.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}txbxContent'):
                    words = [node.text or '' for node in textbox.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')]
                    extra = ' '.join(words).strip()
                    if extra and extra not in text: text += '\n〔文本框文字〕\n' + extra
        elif suffix == ".xlsx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 200 * 1024 * 1024:
                    raise ValueError("Excel 解压后超过 200 MiB，请拆分后上传。")
            try:
                md = MarkItDown()
                result = md.convert_stream(io.BytesIO(content), file_extension=".xlsx")
                text = result.text_content
            except Exception:
                text = _extract_xlsx(content)
        elif suffix == ".pptx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 200 * 1024 * 1024:
                    raise ValueError("PPT 解压后超过 200 MiB，请拆分后上传。")
            md = MarkItDown()
            result = md.convert_stream(io.BytesIO(content), file_extension=".pptx")
            text = result.text_content
        else:
            raise ValueError("不支持的文件格式。")
    except UnicodeDecodeError:
        raise ValueError("文本编码无法识别，请另存为 UTF-8 后上传。") from None
    except ValueError:
        raise
    except Exception:
        raise ValueError("文件无法解析，请检查格式、文件完整性或密码保护。") from None
    text = text.replace("\x00", "").strip()
    if not text and not allow_empty:
        if suffix == ".pdf":
            raise ValueError("未提取到文字；扫描版 PDF 需先进行文字识别。")
        raise ValueError("未提取到有效文字，请检查文件内容。")
    if len(text) > (CHUNK_SIZE - CHUNK_OVERLAP) * MAX_CHUNKS:
        raise ValueError("文件文字超过处理上限，请拆分后上传。")
    return suffix, text


def _office_images(suffix, content):
    import posixpath
    import re
    from xml.etree import ElementTree as ET
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        names = set(archive.namelist())
        if suffix == '.docx':
            parts = [('word/document.xml', 'Word 正文')]
        elif suffix == '.pptx':
            slides = sorted((name for name in names if re.fullmatch(r'ppt/slides/slide\d+\.xml', name)), key=lambda name: int(re.search(r'slide(\d+)', name)[1]))
            parts = [(name, 'PPT 第 ' + re.search(r'slide(\d+)', name)[1] + ' 页') for name in slides]
        else:
            parts = []
            workbook = ET.fromstring(archive.read('xl/workbook.xml'))
            relationships = _relationships(archive, 'xl/workbook.xml')
            for sheet in workbook.iter():
                if sheet.tag.endswith('}sheet'):
                    path = relationships.get(sheet.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'))
                    if not path: continue
                    sheet_root = ET.fromstring(archive.read(path))
                    sheet_rels = _relationships(archive, path)
                    for drawing in sheet_root.iter():
                        if drawing.tag.endswith('}drawing'):
                            drawing_path = sheet_rels.get(drawing.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'))
                            if drawing_path: parts.append((drawing_path, 'Excel 工作表“' + sheet.get('name', '') + '”'))
        for path, label in parts:
            if path not in names: continue
            root = ET.fromstring(archive.read(path)); relationships = _relationships(archive, path)
            count = 0
            containers = [(root, label)]
            if suffix == '.docx':
                containers = [(paragraph, label + ' 第 ' + str(number) + ' 段') for number, paragraph in enumerate(root.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'), 1)]
            elif suffix == '.xlsx':
                from openpyxl.utils import get_column_letter
                containers = []
                for anchor in root:
                    source = anchor.find('{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}from')
                    position = ''
                    if source is not None:
                        column = source.find('{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}col')
                        row = source.find('{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}row')
                        if column is not None and row is not None:
                            position = ' 单元格 ' + get_column_letter(int(column.text) + 1) + str(int(row.text) + 1)
                    containers.append((anchor, label + position))
            for container, location in containers:
                for element in container.iter():
                    if not element.tag.endswith('}blip'): continue
                    target = relationships.get(element.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'))
                    if target and target in names:
                        count += 1
                        yield location + ' 图片 ' + str(count), archive.read(target)


def _relationships(archive, path):
    import posixpath
    from xml.etree import ElementTree as ET
    parent, filename = posixpath.split(path)
    relation_path = posixpath.join(parent, '_rels', filename + '.rels')
    if relation_path not in archive.namelist(): return {}
    result = {}
    for item in ET.fromstring(archive.read(relation_path)):
        if item.get('TargetMode') == 'External': continue
        target = item.get('Target', '')
        result[item.get('Id')] = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(parent, target))
    return result


def extract_document(filename, content, recognizer=None):
    import hashlib
    import time
    from .ocr import recognize_image, render_pdf_page, EXTRACTION_VERSION
    suffix, native = extract_text(filename, content, allow_empty=True)
    recognizer = recognizer or recognize_image
    info = {'version': EXTRACTION_VERSION, 'state': 'not_needed', 'images': 0, 'recognized': 0, 'warnings': []}
    blocks = [native] if native else []
    cache = {}; started = time.monotonic()
    def append_ocr(label, image):
        info['images'] += 1
        if info['images'] > 200 or time.monotonic() - started > 600:
            info['warnings'].append('OCR 达到处理上限，请拆分文档后重试。'); return
        try:
            digest = hashlib.sha256(image).hexdigest()
            if digest not in cache: cache[digest] = recognizer(image)
            text = cache[digest].strip()
            if text:
                blocks.append('〔OCR · ' + label + '〕\n' + text)
                info['recognized'] += 1
        except (ValueError, OSError) as error:
            info['warnings'].append(label + '：' + str(error))
    if suffix == '.pdf':
        reader = PdfReader(io.BytesIO(content))
        for index, page in enumerate(reader.pages):
            page_text = page.extract_text() or ''
            if not page_text.strip() and page.get_contents() is not None:
                append_ocr('PDF 第 ' + str(index + 1) + ' 页', render_pdf_page(content, index))
            elif page_text.strip():
                for number, image in enumerate(page.images, 1):
                    append_ocr('PDF 第 ' + str(index + 1) + ' 页 图片 ' + str(number), image.data)
    elif suffix in {'.docx', '.xlsx', '.pptx'}:
        for label, image in _office_images(suffix, content): append_ocr(label, image)
    info['warnings'] = list(dict.fromkeys(info['warnings']))[:50]
    info['state'] = 'partial' if info['warnings'] else 'complete' if info['images'] else 'not_needed'
    text = '\n\n'.join(blocks).replace('\x00', '').strip()
    if not text:
        if info['warnings']: raise ValueError('未提取到文字，' + info['warnings'][0])
        raise ValueError('未提取到有效文字，扫描版 PDF 或图片需检查 OCR 识别结果。')
    if len(text) > (CHUNK_SIZE - CHUNK_OVERLAP) * MAX_CHUNKS:
        raise ValueError('文字和 OCR 结果超过处理上限，请拆分后上传。')
    return suffix, text, info
