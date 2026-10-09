"""仅生成允许的内容标签；不执行文档 HTML、脚本或外部资源。"""
import csv
import io
from html import escape
from html.parser import HTMLParser
import markdown
import openpyxl
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation


class SafeMarkup(HTMLParser):
    allowed = {'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'p', 'br', 'hr', 'pre', 'code', 'strong', 'em', 'blockquote', 'ul', 'ol', 'li', 'table', 'thead', 'tbody', 'tr', 'td', 'th', 'del'}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag in self.allowed:
            self.parts.append('<' + tag + '>')

    def handle_endtag(self, tag):
        if tag in self.allowed:
            self.parts.append('</' + tag + '>')

    def handle_data(self, data):
        self.parts.append(escape(data))


def table_html(rows):
    return '<table><tbody>' + ''.join('<tr>' + ''.join('<td>' + escape(str(cell) if cell is not None else '') + '</td>' for cell in row) + '</tr>' for row in rows) + '</tbody></table>'


def render_preview(document, path):
    suffix = document['suffix']
    if suffix == '.txt':
        return '<pre>' + escape(document['text']) + '</pre>'
    if suffix == '.md':
        parser = SafeMarkup()
        parser.feed(markdown.markdown(document['text'], extensions=['tables', 'fenced_code']))
        return ''.join(parser.parts)
    if suffix == '.csv':
        raw = path.read_bytes()
        for encoding in ('utf-8-sig', 'gb18030', 'gbk'):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        return table_html(csv.reader(io.StringIO(text)))
    if suffix == '.docx':
        doc = Document(path)
        blocks = []
        for element in doc.element.body.iterchildren():
            if element.tag.endswith('}p'):
                paragraph = Paragraph(element, doc)
                tag = 'h2' if paragraph.style.name.startswith('Heading') else 'p'
                blocks.append(f'<{tag}>' + escape(paragraph.text) + f'</{tag}>')
            elif element.tag.endswith('}tbl'):
                table = Table(element, doc)
                blocks.append(table_html([[cell.text for cell in row.cells] for row in table.rows]))
        return ''.join(blocks)
    if suffix == '.xlsx':
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            blocks = []
            for sheet in workbook:
                rows = (row for row in sheet.iter_rows(values_only=True) if any(cell is not None for cell in row))
                blocks.append('<h2>' + escape(sheet.title) + '</h2>' + table_html(rows))
            return ''.join(blocks)
        finally:
            workbook.close()
    if suffix == '.pptx':
        slides = Presentation(path)
        blocks = []
        for number, slide in enumerate(slides.slides, 1):
            parts = ['<section class="slide-preview"><h2>第 ' + str(number) + ' 页</h2>']
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.extend('<p>' + escape(p.text) + '</p>' for p in shape.text_frame.paragraphs)
                if shape.has_table:
                    parts.append(table_html([[cell.text for cell in row.cells] for row in shape.table.rows]))
            parts.append('</section>')
            blocks.extend(parts)
        return ''.join(blocks)
    raise ValueError('该文件格式无法预览。')
